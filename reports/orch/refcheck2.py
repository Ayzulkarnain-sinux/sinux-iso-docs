#!/usr/bin/env python3
"""ISO-W9 (LEAD G-2 butir 2): cek rujukan "butir/bagian/langkah/prosedur N.M" dan "K-nn" di semua docx out/ dengan
keluaran penuh per rujukan. Nomor heading Word diturunkan dari urutan heading bernomor (numPr), bukan dicari sebagai
teks. Status per rujukan:
  ADA              tujuan ada dan nomornya tampak (heading 2 ke-M di bagian N, paragraf "N.M"/"N.", baris langkah N)
  ADA_TANPA_NOMOR  hanya ada sebagai paragraf ke-M tanpa nomor tampak (pembaca harus menghitung)
  TIDAK_ADA        nomor itu tidak ada di dokumen tujuan
  DOK_TIDAK_ADA    kode dokumen tidak ada di set DONE
  PERLU_CEK_MAKNA  rujukan ke uraian bernomor di dalam sel tabel (mis. "JA-12 butir 6") atau ke xlsx; diputus pemeriksa makna
Kontrol positif: SNX-OPS-FRM-10 versi commit 4a40352 harus menembak "SNX-OPS-SOP-04 Bagian 7.5" (TIDAK_ADA).

Pakai (root repo): python3 -I reports/orch/refcheck2.py <tag>
  -> reports/w9/refcheck_<tag>.tsv dan .json; ringkasan ke stdout; keluar 1 bila ada TIDAK_ADA/DOK_TIDAK_ADA
     atau kontrol positif gagal."""
import collections, io, json, os, re, subprocess, sys
import docx
from docx.oxml.ns import qn

L = json.load(open('ledger.json'))
DOCS = [d for d in L['docs'] if d.get('out_path') and d['state'] == 'DONE']
BY_NUM = {d['number']: d for d in DOCS}
NUM_OF = {d['key']: d['number'] for d in DOCS}
CODE = r'SNX-[A-Z]+-[A-Z]+-\d\d'
NUM = r'[A-Z]?\d+(?:\.\d+)*'
LOC = re.compile(r'\b(butir|bagian|langkah|prosedur)\s+(' + NUM + r')((?:\s*(?:,|dan|sampai|hingga|atau)\s*(?:butir\s+|bagian\s+|langkah\s+)?' + NUM + r'\b)*)', re.I)
KREF = re.compile(r'\b(K-\d\d)\b')
HINTS = r'(Ketentuan|Petunjuk Pengisian|Cara Mengisi|Daftar Periksa|Tujuan|Peran|Rekaman)'
CTRL_COMMIT = '4a40352'


def ptext(el):
    return ''.join(t.text or '' for t in el.iter(qn('w:t')))


def parse(blob):
    """-> (unit list [(h1, h2, teks, kode_baris)], struktur)."""
    d = docx.Document(io.BytesIO(blob))
    body = d.element.body
    sdts = [i for i, e in enumerate(body) if e.tag == qn('w:sdt')]
    st = {'secs': {}, 'subs': {}, 'items': {}, 'ordinal': collections.defaultdict(list), 'litems': collections.defaultdict(dict),
          'steps': collections.defaultdict(dict), 'cells': collections.defaultdict(list), 'cellnum': collections.defaultdict(set),
          'tnum': {}}
    out, h1, h2 = [], 0, 0
    for el in list(body)[(sdts[0] + 1 if sdts else 0):]:
        if el.tag == qn('w:p'):
            ps = el.find('.//' + qn('w:pStyle'))
            style = ps.get(qn('w:val')) if ps is not None else ''
            numbered = el.find('.//' + qn('w:numPr')) is not None
            txt = ptext(el).strip()
            if style == 'Heading1' and numbered:
                h1, h2 = h1 + 1, 0
                st['secs'][h1] = txt
                continue
            if style == 'Heading2' and numbered:
                h2 += 1
                st['subs'][(h1, h2)] = txt
                continue
            if not txt or not h1:
                continue
            m, m1 = re.match(r'^(\d+)\.(\d+)\b', txt), re.match(r'^(\d+)\.\s', txt)
            if m:
                st['items'][(int(m.group(1)), int(m.group(2)))] = txt
            elif m1:
                st['litems'][h1][int(m1.group(1))] = txt
            else:
                st['ordinal'][h1].append(txt)
            out.append((h1, h2, txt, None, 'p'))
        elif el.tag == qn('w:tbl') and h1:
            for tr in el.iter(qn('w:tr')):
                tcs = list(tr.iter(qn('w:tc')))
                cells = [ptext(tc).strip() for tc in tcs]
                if not cells:
                    continue
                st['cells'][h1].append(cells)
                ms = re.match(r'^(\d+)(?:\.\s|$)', cells[0])
                if ms:  # langkah prosedur ("1. Kirim ...") atau baris tugas bernomor ("1")
                    st['steps'][(h1, h2)][int(ms.group(1))] = ' | '.join(cells)
                mt = re.fullmatch(r'(\d+)\.(\d+)', cells[0])
                if mt:  # nomor butir checklist di sel pertama tabel (mis. 2.4)
                    st['tnum'][(int(mt.group(1)), int(mt.group(2)))] = (h1, ' | '.join(cells))
                for tc in tcs:
                    for p in tc.iter(qn('w:p')):
                        mm = re.match(r'^(\d+)\.\s', ptext(p).strip())
                        if mm:
                            st['cellnum'][h1].add(int(mm.group(1)))
                rc = re.match(r'^(' + CODE + r')\b', cells[0])
                out.append((h1, h2, ' | '.join(cells), None, 'row'))  # hanya untuk kerangka (dump), dilewati pemeriksa
                for i, c in enumerate(cells):
                    if c:
                        out.append((h1, h2, c, rc.group(1) if rc and i else None, 't'))
    return out, st


def resolve(st, kw, loc, hint, ctx, unit_txt):
    kw = kw.lower()
    if re.match(r'^[A-Z]', loc):  # ID baris tabel (C3.2, A6, U14)
        for sec, rows in st['cells'].items():
            for r in rows:
                if r[0] == loc or r[0].startswith((loc + ' ', loc + '.')):
                    return 'ADA', f"bagian {sec} {st['secs'].get(sec, '')}: {' | '.join(r)[:110]}"
        return 'TIDAK_ADA', 'ID baris tidak ditemukan di tabel mana pun'
    parts = [int(x) for x in loc.split('.')]
    n = parts[0]
    if len(parts) == 1 and kw in ('butir', 'langkah'):
        explicit = isinstance(ctx, tuple) and len(ctx) == 3
        c2 = ctx[:2] if ctx else None
        if c2 and c2 in st['subs']:  # "6.4 langkah 2", "langkah 2 butir 6.4", "Bagian 6.3 butir 4", langkah di tabel heading 2 sendiri
            s = st['steps'].get(c2, {})
            if n in s:
                return 'ADA', f"{c2[0]}.{c2[1]} {st['subs'][c2]} > langkah {n}: {s[n][:90]}"
            if explicit:
                return 'TIDAK_ADA', f"{c2[0]}.{c2[1]} {st['subs'][c2]} punya langkah {', '.join(map(str, sorted(s))) or 'tanpa nomor'}"
        if re.search(r'\bJA-\d\d\b|sheet|lembar', unit_txt, re.I):
            return 'PERLU_CEK_MAKNA', 'uraian bernomor di sel tabel (JA) atau sheet xlsx'
        cand = [sec for sec, li in st['litems'].items() if n in li]
        hinted = [sec for sec in cand if hint and hint.lower() in st['secs'].get(sec, '').lower()]
        pick = hinted or cand
        if pick:
            sec = pick[0]
            amb = '' if len(pick) == 1 else f" (juga di bagian {', '.join(map(str, pick[1:]))})"
            return 'ADA', f"bagian {sec} {st['secs'].get(sec, '')} > {st['litems'][sec][n][:100]}{amb}"
        steps_any = [k for k, s in st['steps'].items() if n in s]
        if kw == 'langkah' and steps_any:
            return 'ADA', f"langkah {n} ada di {', '.join(f'{a}.{b}' for a, b in steps_any)} (prosedur tidak disebut)"
        if any(n in v for v in st['cellnum'].values()):
            return 'PERLU_CEK_MAKNA', f"butir {n} hanya ada sebagai uraian bernomor di sel tabel"
        return 'TIDAK_ADA', f"tidak ada butir bernomor tampak '{n}.'{' di bagian ' + hint if hint else ''}"
    if len(parts) == 2 and tuple(parts) in st['tnum'] and kw == 'butir':
        sec, txt = st['tnum'][tuple(parts)]
        return 'ADA', f"bagian {sec} {st['secs'].get(sec, '')} > baris tabel {loc}: {txt[:100]}"
    if n not in st['secs']:
        return 'TIDAK_ADA', f"dokumen punya bagian 1 sampai {max(st['secs'] or [0])}"
    title = f"bagian {n} {st['secs'][n]}"
    if len(parts) == 1:
        return 'ADA', title
    m = parts[1]
    if (n, m) in st['subs']:
        return 'ADA', f"{title} > {n}.{m} {st['subs'][(n, m)]}"
    if (n, m) in st['items']:
        return 'ADA', f"{title} > {st['items'][(n, m)][:110]}"
    if len(st['ordinal'][n]) >= m and not any(k[0] == n for k in st['items']) and not any(k[0] == n for k in st['subs']):
        return 'ADA_TANPA_NOMOR', f"{title} > paragraf ke-{m} tanpa nomor tampak: {st['ordinal'][n][m - 1][:100]}"
    have = sorted([k[1] for k in st['subs'] if k[0] == n] + [k[1] for k in st['items'] if k[0] == n])
    return 'TIDAK_ADA', f"{title}; nomor tampak di bagian ini: {', '.join(f'{n}.{x}' for x in have) or 'tidak ada'}"


def kresolve(st, k):
    for sec, rows in st['cells'].items():
        for r in rows:
            if r[0] == k:
                return 'ADA', f"bagian {sec} {st['secs'].get(sec, '')}: {' | '.join(r)[:110]}"
    return 'TIDAK_ADA', 'baris K tidak ada'


def parse_xlsx(blob):
    """Unit teks sel xlsx (asal rujukan); struktur kosong (penomoran butir xlsx bebas)."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(blob))
    out = [(0, 0, f'{ws.title}!{c.coordinate}: {c.value}', None, 't') for ws in wb.worksheets for r in ws.iter_rows() for c in r
           if isinstance(c.value, str) and not c.value.startswith('=')]
    st = {'secs': {}, 'subs': {}, 'items': {}, 'ordinal': collections.defaultdict(list), 'litems': collections.defaultdict(dict),
          'steps': collections.defaultdict(dict), 'cells': collections.defaultdict(list), 'cellnum': collections.defaultdict(set), 'tnum': {}}
    return out, st


def refs_of(key, blob, structs):
    us, own = parse_xlsx(blob) if BY_NUM[NUM_OF[key]]['out_path'].endswith('.xlsx') else parse(blob)
    self_num = NUM_OF[key]
    rows = []
    for h1, h2, t, rowcode, kind_u in us:
        if kind_u == 'row':
            continue
        codes = [(m.start(), m.end(), m.group(0)) for m in re.finditer(CODE, t)]
        for m in LOC.finditer(t):
            kw = m.group(1)
            # tech_fix ISO-W10: daftar campuran ("butir 4.7 dan bagian 11") memakai kata kunci tiap anggota, bukan kata kunci kepala.
            locs = [(kw, m.group(2))]
            for kk, nn in re.findall(r'(?:\b(butir|bagian|langkah)\s+)?(' + NUM + r')', m.group(3) or '', re.I):
                locs.append((kk or locs[-1][0], nn))
            hm = re.search(HINTS + r'\s*$', t[max(0, m.start() - 25):m.start()])
            hint = hm.group(1) if hm else ''
            before = [c for c in codes if c[1] <= m.start() and m.start() - c[1] <= 45 and not re.search(r'[.;:]\s|' + CODE, t[c[1]:m.start()])]
            after = [c for c in codes if c[0] >= m.end() and c[0] - m.end() <= 12 and re.match(r'^\s*(?:di|pada|dalam|dari)?\s*\(?$', t[m.end():c[0]])]
            tgt = before[-1][2] if before else after[0][2] if after else rowcode
            # konteks prosedur untuk langkah/butir satu angka: "6.4 langkah 2", "langkah 2 butir 6.4", "Bagian 6.3 butir 4"
            ctx = None
            pre = re.search(r'\b(\d+)\.(\d+)\s*$', t[max(0, m.start() - 12):m.start()])
            post = re.match(r'^\s*(?:butir|prosedur|bagian)?\s*(\d+)\.(\d+)\b', t[m.end():m.end() + 20])
            if pre:
                ctx = (int(pre.group(1)), int(pre.group(2)), 'eksplisit')
            elif post and kw.lower() == 'langkah':
                ctx = (int(post.group(1)), int(post.group(2)), 'eksplisit')
            elif not tgt and not hint and h2 and kw.lower() in ('langkah', 'butir'):
                ctx = (h1, h2)  # implisit: langkah atau uraian bernomor di tabel heading 2 tempat kalimat itu berada
            for kwi, loc in locs:
                rows.append(make_row(key, self_num, h1, t, tgt, kwi, loc, hint, ctx, structs, own))
        for m in KREF.finditer(t):
            rows.append(make_row(key, self_num, h1, t, 'SNX-HR-ORG-01', 'K', m.group(1), '', None, structs, own))
    return rows


def make_row(key, self_num, sec, t, tgt, kw, loc, hint, ctx, structs, own):
    kind = 'silang' if tgt and tgt != self_num else 'internal'
    num = tgt or self_num
    if num not in BY_NUM:
        status, title = 'DOK_TIDAK_ADA', 'nomor tidak ada di set dokumen DONE'
    elif not BY_NUM[num]['out_path'].endswith('.docx'):
        status, title = 'PERLU_CEK_MAKNA', 'tujuan xlsx: penomoran butir bebas'
    elif kind == 'internal' and not own['secs']:
        status, title = 'PERLU_CEK_MAKNA', 'rujukan internal di xlsx: penomoran butir bebas'
    else:
        st = own if kind == 'internal' else structs[num]
        status, title = kresolve(st, loc) if kw == 'K' else resolve(st, kw, loc, hint, ctx, t)
    lok = loc if kw == 'K' else ' '.join(x for x in (hint, kw, loc) if x)
    return {'asal': key, 'bagian_asal': sec, 'jenis': kind, 'tujuan': num, 'lokasi': lok, 'status': status,
            'judul_tujuan': title, 'kalimat': t[:400]}


def main():
    tag = sys.argv[1]
    blobs = {d['key']: open(d['out_path'], 'rb').read() for d in DOCS}
    structs = {NUM_OF[k]: parse(b)[1] for k, b in blobs.items() if BY_NUM[NUM_OF[k]]['out_path'].endswith('.docx')}
    rows = [r for k, b in blobs.items() for r in refs_of(k, b, structs)]
    fp = BY_NUM['SNX-OPS-FRM-10']['out_path']
    old = subprocess.run(['git', 'show', f'{CTRL_COMMIT}:{fp}'], capture_output=True).stdout
    ctrl = [r for r in refs_of('OPS-FRM-10', old, structs) if r['tujuan'] == 'SNX-OPS-SOP-04' and r['lokasi'].endswith(' 7.5')]
    ctrl_ok = bool(ctrl) and all(r['status'] == 'TIDAK_ADA' for r in ctrl)
    os.makedirs('reports/w9', exist_ok=True)
    cols = ('asal', 'bagian_asal', 'jenis', 'tujuan', 'lokasi', 'status', 'judul_tujuan', 'kalimat')
    with open(f'reports/w9/refcheck_{tag}.tsv', 'w', encoding='utf-8') as f:
        f.write('\t'.join(cols) + '\n')
        for r in rows:
            f.write('\t'.join(str(r[c]).replace('\t', ' ').replace('\n', ' ') for c in cols) + '\n')
    json.dump({'rows': rows, 'control': ctrl, 'control_ok': ctrl_ok}, open(f'reports/w9/refcheck_{tag}.json', 'w'), ensure_ascii=False, indent=1)
    c = collections.Counter((r['jenis'], r['status']) for r in rows)
    uniq = {(r['asal'], r['tujuan'], r['lokasi']) for r in rows if r['jenis'] == 'silang'}
    print(f"KONTROL POSITIF FRM-10 @{CTRL_COMMIT} 'SNX-OPS-SOP-04 Bagian 7.5': {len(ctrl)} kemunculan, "
          f"{'OK menembak TIDAK_ADA' if ctrl_ok else 'GAGAL'}")
    print(f"rujukan: {len(rows)} kemunculan; silang unik (asal, tujuan, lokasi): {len(uniq)}; berkas dibaca: {len(blobs)} (docx {len(structs)}, xlsx {len(blobs) - len(structs)})")
    for k in sorted(c):
        print(f"  {k[0]:8s} {k[1]:16s} {c[k]}")
    # Adjudikasi pemeriksa makna (reports/w9/refsem*.json): rujukan mesin TIDAK_ADA/DOK_TIDAK_ADA/ADA_TANPA_NOMOR yang
    # diputus OK (rujukan ke dokumen asal sendiri) atau ke dokumen tim di luar set ditandai DIPUTUS_MAKNA dengan id buktinya.
    adj = {}
    import glob
    for f in sorted(glob.glob('reports/w9/refsem*.json')):
        for r in json.load(open(f)).get('rows', []):
            s = r.get('sem') or {}
            ok = s.get('verdict') == 'OK' or (s.get('verdict') == 'SALAH_TAFSIR_MESIN' and s.get('verdict_ulang') in ('OK', 'TIDAK_DAPAT_DICEK'))
            if ok:
                adj[(r['asal'], r['tujuan'], r['lokasi'], r['kalimat'][:400])] = r['id']
    for r in rows:
        k = (r['asal'], r['tujuan'], r['lokasi'], r['kalimat'][:400])
        if r['status'] in ('TIDAK_ADA', 'DOK_TIDAK_ADA', 'ADA_TANPA_NOMOR') and k in adj:
            r['status'] = 'DIPUTUS_MAKNA'
            r['judul_tujuan'] = f"diputus pemeriksa makna {adj[k]}; " + r['judul_tujuan']
    json.dump({'rows': rows, 'control': ctrl, 'control_ok': ctrl_ok}, open(f'reports/w9/refcheck_{tag}.json', 'w'), ensure_ascii=False, indent=1)
    with open(f'reports/w9/refcheck_{tag}.tsv', 'w', encoding='utf-8') as f:
        f.write('\t'.join(cols) + '\n')
        for r in rows:
            f.write('\t'.join(str(r[c]).replace('\t', ' ').replace('\n', ' ') for c in cols) + '\n')
    c = collections.Counter((r['jenis'], r['status']) for r in rows)
    print('sesudah adjudikasi:')
    for k in sorted(c):
        print(f"  {k[0]:8s} {k[1]:16s} {c[k]}")
    bad = [r for r in rows if r['status'] in ('TIDAK_ADA', 'DOK_TIDAK_ADA')]
    for r in bad:
        print(f"RUSAK {r['asal']} -> {r['tujuan']} {r['lokasi']} [{r['status']}] {r['judul_tujuan'][:80]} :: {r['kalimat'][:130]}")
    sys.exit(1 if bad or not ctrl_ok else 0)


if __name__ == '__main__' and sys.argv[1] != 'dump':
    main()


def dump(outdir):
    """Kerangka bernomor tiap docx DONE (bahan pemeriksa makna): bagian, heading 2, butir tampak, paragraf tanpa nomor, langkah, baris tabel."""
    os.makedirs(outdir, exist_ok=True)
    for d in DOCS:
        if not d['out_path'].endswith('.docx'):
            continue
        us, st = parse(open(d['out_path'], 'rb').read())
        lines = [f"# {d['number']} {d['out_path'].split('/')[-1]}"]
        cur, k = None, 0
        for h1, h2, t, rc, ku in us:
            if ku == 't':
                continue
            if (h1, h2) != cur:
                cur, k = (h1, h2), 0
                lines.append(f"\n## bagian {h1} {st['secs'].get(h1, '')}" if not h2 else f"\n### {h1}.{h2} {st['subs'].get((h1, h2), '')}")
            if ku == 'p' and not re.match(r'^\d+\.', t):
                k += 1
                lines.append(f"[paragraf tanpa nomor ke-{k} di bagian {h1}] {t}")
            else:
                lines.append(('[baris tabel] ' if ku == 'row' else '') + t)
        open(f"{outdir}/{d['key']}.txt", 'w', encoding='utf-8').write('\n'.join(lines))
    print('ok', outdir)


if __name__ == '__main__' and len(sys.argv) > 2 and sys.argv[1] == 'dump':
    dump(sys.argv[2])
