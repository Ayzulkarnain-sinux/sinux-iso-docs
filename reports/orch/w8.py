#!/usr/bin/env python3
"""W8: bangun out/00_Index/Master_List_COS.xlsx dari ledger, worklist, checklist, notes, qa.
Pakai (dari root repo): python3 -I reports/orch/w8.py  -> cetak ringkasan JSON untuk Laporan_Run.md."""
import glob, json, os, re
from collections import Counter, defaultdict
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

L = json.load(open('ledger.json'))
W = json.load(open('config/worklist.json'))['items']
C = json.load(open('config/checklist.json'))
rows = {r['row']: r for r in (C['rows'] if isinstance(C, dict) else C)}
DEPT = {'HR': 'Human Resource', 'FAT': 'Finance, Accounting & Tax', 'LEG': 'Legal', 'COM': 'Commercial',
        'MKT': 'Marketing', 'OPS': 'Operation', 'ISMS': 'ISMS'}
GROUPS = [('owner', r'\bowner\b|Amirull'), ('Samuel', r'Samuel|Direktur Utama'), ('Aziz', r'Aziz|HR Manager'),
          ('Fani', r'Fani|VP Tech'), ('FAT', r'\bFAT\b|Finance'), ('Legal', r'Legal'), ('Laras', r'Laras|VP Commercial'),
          ('konsultan', r'konsultan'), ('LEAD', r'\bLEAD\b')]


def section(text, name):
    m = re.search(rf'^## {re.escape(name)}[^\n]*\n(.*?)(?=^## |\Z)', text, re.S | re.M)
    return m.group(1) if m else ''


def items(sec):
    """Butir dari tabel markdown (baris judul dan pemisah dilewati) atau daftar berpoin."""
    out, in_table = [], False
    for line in sec.splitlines():
        s = line.strip()
        if s.startswith('|'):
            cells = [c.strip() for c in s.strip('|').split('|')]
            if not in_table:  # baris judul tabel
                in_table = True
                continue
            if all(re.fullmatch(r':?-{2,}:?', c) for c in cells if c):
                continue
            body = [c for c in cells if c and not re.fullmatch(r'[A-Z]{0,3}-?\d+[a-z]?', c)]
            if body:
                out.append((' | '.join(body[:-1]) or body[0], body[-1]))
            continue
        in_table = False
        m = re.match(r'(?:[-*]|\d+[.)])\s+(.*)', s)
        if m:
            t = m.group(1)
            p = re.search(r'[Pp]emutus[^:]*:\s*(.+?)\.?$', t) or re.search(r'\(([^()]*(?:owner|Samuel|Aziz|Fani|FAT|Legal|Laras|konsultan)[^()]*)\)\s*\.?$', t)
            out.append((t, p.group(1) if p else ''))
    return out


def groups_of(who, text):
    g = [n for n, rx in GROUPS if re.search(rx, who or '')] or [n for n, rx in GROUPS if re.search(rx, text)][:1]
    return g or ['lainnya']


def verify_label(d):
    v = d.get('verify_verdict')
    rounds = len(glob.glob(f"qa/{d['key']}.verify.r*.json")) + (1 if os.path.exists(f"qa/{d['key']}.verify.json") else 0)
    if v == 'PASS':
        return f'PASS (pemeriksaan independen, putaran {rounds})'
    if d['state'] == 'DONE':
        return 'Tidak disampel (aturan sampel W5/W6)'
    return v or ''


def pdf_label(key):
    p = f'qa/{key}.pdfcheck.json'
    if not os.path.exists(p):
        return ''
    try:
        return json.load(open(p)).get('verdict', '')
    except Exception:
        return 'tidak terbaca'


DEFAULT = {'HR': 'Aziz', 'FAT': 'FAT', 'LEG': 'Legal', 'COM': 'owner', 'MKT': 'owner', 'OPS': 'Fani', 'ISMS': 'Fani'}


def doc_usulan(path):
    """(lokasi, teks) untuk tiap paragraf atau sel yang memuat label (USULAN dari dokumen jadi (teks sudah lolos lint)."""
    out = []
    if path.endswith('.xlsx'):
        wb = openpyxl.load_workbook(path)
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and '(USULAN' in c.value:
                        out.append((f'{ws.title}!{c.coordinate}', c.value))
        return out
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = docx.Document(path)
    head, seen = '', set()
    for el in d.element.body.iterchildren():
        if el.tag.endswith('}p'):
            para = Paragraph(el, d)
            if para.style.name.lower().startswith('heading') and para.text.strip():
                head = para.text.strip()
            elif '(USULAN' in para.text:
                out.append((head, para.text.strip()))
        elif el.tag.endswith('}tbl'):
            for r in Table(el, d).rows:
                t = ' | '.join(dict.fromkeys(c.text.strip() for c in r.cells if c.text.strip()))
                if '(USULAN' in t and t not in seen:
                    seen.add(t)
                    out.append((head, t))
    return out


def words(t):
    return set(re.findall(r'[a-zA-Z]{4,}', t.lower()))


def match_group(text, cands, dept):
    if 'konfirmasi konsultan' in text:
        return 'konsultan'
    w = words(text)
    best = max(cands, key=lambda c: len(w & c[0]) / (len(w | c[0]) or 1), default=None)
    if best and len(w & best[0]) / (len(w | best[0]) or 1) >= 0.2:
        return best[1]
    return DEFAULT.get(dept, 'lainnya')


wl = {i: w for i, w in enumerate(W)}
docu = []
master, usul, opens = [], [], []
for d in sorted(L['docs'], key=lambda x: (x['checklist_row'] is None, x['checklist_row'] or 0, x['wl_index'])):
    w = wl.get(d['wl_index'], {})
    r = rows.get(d['checklist_row'], {})
    title = r.get('title') or w.get('checklist_title') or ''
    srcs = '; '.join(f"{s.get('title') or s['id']} ({s['id']}, {s.get('modifiedTime', 'modifiedTime tidak diukur')})" for s in d.get('sources') or [])
    notes = open(f"notes/{d['key']}.md", encoding='utf-8').read() if d.get('key') and os.path.exists(f"notes/{d['key']}.md") else ''
    u, o = items(section(notes, 'USULAN')), items(section(notes, 'Open items'))
    for t, who in u:
        usul.append([d['number'], d['key'], t, who, ', '.join(groups_of(who, t))])
    for t, who in o:
        opens.append([d['number'], d['key'], t, who, ', '.join(groups_of(who, t))])
    if d['state'] == 'DONE' and d.get('out_path') and os.path.exists(d['out_path']):
        cands = [(words(t), groups_of(who, t)[0]) for t, who in u]
        for loc, t in doc_usulan(d['out_path']):
            docu.append([d['number'], DEPT.get(d['dept'], d['dept']), loc, t, match_group(t, cands, d['dept'])])
    state = d['state']
    ket = {'HOLD': 'Menunggu berkas tim (T3); ' + (w.get('notes') or ''), 'SKIP': w.get('notes') or ''}.get(state, '')
    master.append([d['checklist_row'] or '(tidak ada di checklist live)', DEPT.get(d['dept'], d['dept']), title, d['action'],
                   d.get('wave'), d.get('number') or '', os.path.basename(d.get('out_path') or ''), state,
                   verify_label(d), pdf_label(d['key']) if d.get('key') else '', d.get('usulan_count') or (len(u) if u else ''),
                   len(o) if o else '', ket, srcs])

wb = openpyxl.Workbook()
HEAD = PatternFill('solid', fgColor='1F3864')
def sheet(ws, header, data, widths):
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color='FFFFFF'), HEAD
        c.alignment = Alignment(wrap_text=True, vertical='top')
    for row in data:
        ws.append(row)
    for i, wdt in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

ws = wb.active; ws.title = 'Master List'
sheet(ws, ['Baris checklist', 'Departemen', 'Judul checklist', 'Aksi', 'Wave', 'Nomor', 'Nama berkas (out/)', 'State',
           'Verifikasi isi', 'Cek PDF', 'Label (USULAN)', 'Open items', 'Keterangan', 'Sumber (fileId, modifiedTime)'],
      master, [9, 16, 40, 6, 7, 18, 50, 9, 26, 9, 9, 9, 40, 70])
sheet(wb.create_sheet('USULAN'), ['Nomor', 'Departemen', 'Bagian', 'Kalimat atau baris berlabel (USULAN)', 'Pemutus disarankan'], docu, [18, 16, 30, 100, 14])
with open('reports/USULAN-OPEN-per-pemutus.md', 'w', encoding='utf-8') as f:
    f.write('# USULAN dan open items per pemutus (dari notes penyusun, bahan kerja internal)\n')
    for title, data in (('USULAN', usul), ('Open items', opens)):
        by = defaultdict(list)
        for x in data:
            for g in x[4].split(', '):
                by[g].append(x)
        f.write(f'\n## {title} ({len(data)})\n')
        for g in [n for n, _ in GROUPS] + ['lainnya']:
            if by.get(g):
                f.write(f'\n### {g} ({len(by[g])})\n')
                for x in by[g]:
                    f.write(f"- {x[0]}: {x[2].replace(chr(10), ' ')} (pemutus: {x[3]})\n")
k = wb.create_sheet('Keterangan')
for line in ['Master List COS: satu baris per butir worklist (88), diurutkan menurut baris checklist live.',
             'Sheet USULAN memuat tiap kalimat atau baris berlabel (USULAN) di dokumen jadi; kolom pemutus adalah saran untuk ditinjau owner.',
             'State: DONE = draf selesai dan lolos gerbang; HOLD = menunggu berkas tim (T3); SKIP = SoA (lane konsultan) atau cetakan SOP.',
             'Verifikasi isi: pemeriksa independen membandingkan dokumen dengan sumber dan kriteria checklist; form W5 dan register Legal W6 disampel 1 dari 3 ditambah semua dokumen dengan lebih dari 5 USULAN.',
             'Semua dokumen berstatus Draft (Rev00). Tanda tangan dan tanggal efektif diisi saat pengesahan.']:
    k.append([line])
k.column_dimensions['A'].width = 140
os.makedirs('out/00_Index', exist_ok=True)
wb.save('out/00_Index/Master_List_COS.xlsx')

per = defaultdict(Counter)
for d in L['docs']:
    per[DEPT.get(d['dept'], d['dept'])][d['state']] += 1
gu, go = Counter(x[4] for x in docu), Counter()
for x in opens:
    for g in x[4].split(', '): go[g] += 1
print(json.dumps({'rows': len(master), 'usulan_items_doc': len(docu), 'usulan_items_notes': len(usul), 'open_items': len(opens), 'per_dept': per,
                  'usulan_per_group': gu, 'open_per_group': go,
                  'states': Counter(d['state'] for d in L['docs'])}, ensure_ascii=False, default=dict))
