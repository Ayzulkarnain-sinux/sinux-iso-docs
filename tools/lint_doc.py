#!/usr/bin/env python3
"""Uji mesin untuk docx hasil render (format template + larangan isi).

Pakai: python3 -I lint_doc.py <doc.docx> [sop|generic]
Keluaran: JSON {pass, fails[], warns[]}; exit 0 bila pass, 1 bila gagal.
"""
import json, re, sys
import docx
from docx.oxml.ns import qn

SOP_HEADINGS = ["PURPOSE", "SCOPE", "DEFINITION & ABBREVIATIONS", "POLICY & GENERAL PRINCIPLES",
                "ROLES, RESPONSIBILITIES & AUTHORITIES", "STANDARD OPERATING PROCEDURE", "PROCESS FLOW",
                "SERVICE LEVEL AGREEMENT (SLA)", "EXCEPTION MANAGEMENT", "CONCLUSION",
                "RELATED DOCUMENTS & ATTACHMENTS"]
FORBIDDEN = [
    (r'SKILL_', 'rujukan skill AI internal'), (r'\bDOCTRINE\b', 'rujukan DOCTRINE'),
    (r'\bDEC-20\d\d', 'rujukan DEC'), (r'\bHANDOVER([_-]\w|\s*§)', 'rujukan HANDOVER internal'),
    (r'\b(?!ISO\b)[A-Z]{2,}\s*§\s*\d', 'rujukan section kanon (KODE §n)'),
    (r'\b[A-Z]{2,}_[A-Z][A-Z_]+\b', 'nama berkas internal (HURUF_BESAR_GARIS_BAWAH)'),
    (r'(?i:submit_to_staging|\bpacket staging\b|\bstaging KB\b|\bantrean staging\b|\bjalur staging\b|\bvia staging\b)', 'jalur staging KB'), (r'packet setor', 'packet setor'),
    (r'(?i:\b(Claude|Altair|kb_anchor|relay_get|relay_mark|query_kb)\b)', 'rujukan AI internal'),
    (r'\bCONTOH\b', 'baris CONTOH'), (r'Nexus AI', 'produk pensiun'), (r'SecureEdge', 'produk tidak ada'),
    (r'Imunify360', 'POLRI = ImunifyAV'), (r'Komisaris dan Advisor', 'peran Fani keliru'),
    ('—', 'em-dash'), (r'\(Menjelaskan', 'teks instruksi template'), (r'\(Menampilkan', 'teks instruksi template'),
    (r'\(Detail Process\)', 'teks instruksi template'),
    (r'\[(NAMA|NOMOR|DEPTARTEMEN|DEPARTEMEN|DD|MONTHLY|Prepared|Reviewed|Internal / Confidential)[^\]]*\]', 'placeholder template'),
    (r'\[ISI', 'placeholder [ISI'), (r'\bTBD\b|\bTODO\b|\?\?\?', 'penanda belum selesai'),
    (r'BEGIN (RSA |OPENSSH )?PRIVATE KEY', 'kunci privat'),
    (r'\b(password|passwd|pwd|kata sandi|sandi|token|api[_ ]?key|secret)(\s+\w+)?\s*[:=]\s*(?!SNX-)(?=\S*\d)(?=\S*[A-Za-z])\S{8,}', 'kemungkinan kredensial'),
]

# Gaya bahasa (anti AI slop) dan ejaan baku, tidak peka huruf besar-kecil. Sumber: RUNBOOK §6a.
STYLE = [
    (r'\b(era digital|era modern|dewasa ini)\b', 'pembuka generik'),
    (r'\b(sangat|amat) penting (untuk|bagi)\b', 'penegasan kosong'),
    (r'memegang peran(an)? (yang )?(sangat )?penting', 'penegasan kosong'),
    (r'\b(holistik|sinergi|bersinergi)\b', 'kata kosong'),
    (r'secara komprehensif', 'kata kosong'),
    (r'tidak hanya .{1,80}? tetapi juga', 'paralelisme negatif'),
    (r'(perlu|penting untuk) dicatat bahwa', 'frasa pengisi'),
    (r'dengan demikian,? dapat disimpulkan', 'penutup generik'),
    (r'\b(seamless|robust|leverage|cutting-edge|state-of-the-art|best-in-class|world-class|game[- ]changer|next-level)\b', 'kosakata AI'),
    (r'\b(kelas dunia|terdepan di industri|revolusioner)\b', 'klaim tanpa bukti'),
    (r'\b(semoga bermanfaat|mari kita)\b', 'gaya obrolan'),
    ('[\U0001F300-\U0001FAFF\u2600-\u260F\u2613-\u2712\u2718-\u2793\u27B0-\u27BF]', 'emoji (kotak centang, tanda cek, dan panah boleh)'),
    (r'\bresiko\b', 'ejaan: risiko'), (r'\b(menganalisa|analisa)\b', 'ejaan: analisis'),
    (r'\bpraktek\b', 'ejaan: praktik'), (r'\baktifitas\b', 'ejaan: aktivitas'), (r'\bijin\b', 'ejaan: izin'),
    (r'\bmerubah\b', 'ejaan: mengubah'), (r'\bsekedar\b', 'ejaan: sekadar'), (r'\bobyek\b', 'ejaan: objek'),
    (r'\bstandarisasi\b', 'ejaan: standardisasi'), (r'\bkwalitas\b', 'ejaan: kualitas'),
    (r'\b(dibawah|diatas|didalam|diluar|kedalam)\b', 'ejaan: kata depan dipisah'),
]
# Sisa isi dummy/uji; tidak berlaku untuk berkas uji ber-nomor SNX-UJI-.
DUMMY = r'Nama Penyusun / Jabatan|Nama Pemeriksa / Jabatan|uji render|lorem ipsum|isi dummy'


def all_text(d):
    parts = []
    for el in d.element.body.iter(qn('w:t')):
        parts.append(el.text or '')
    for s in d.sections:
        for hf in (s.header, s.footer):
            for el in hf._element.iter(qn('w:t')):
                parts.append(el.text or '')
    return '\n'.join(parts)


def para_text(el):
    return ''.join(t.text or '' for t in el.iter(qn('w:t'))).strip()


def main():
    path = sys.argv[1]
    mode = sys.argv[2] if len(sys.argv) > 2 else 'sop'
    d = docx.Document(path)
    fails, warns = [], []
    txt = all_text(d)
    for pat, why in FORBIDDEN:
        m = re.search(pat, txt, flags=re.I if 'password' in pat else 0)
        if m:
            fails.append(f'terlarang ({why}): "{m.group(0)[:40]}"')

    for pat, why in STYLE:
        m = re.search(pat, txt, flags=re.I)
        if m:
            fails.append(f'gaya ({why}): "{m.group(0)[:40]}"')
    if not (d.tables and len(d.tables[0].rows) > 2 and 'SNX-UJI-' in ' '.join(c.text for r in d.tables[0].rows for c in r.cells)):
        m = re.search(DUMMY, txt, flags=re.I)
        if m:
            fails.append(f'sisa isi uji/dummy: "{m.group(0)}"')
    for tag in ('w:bookmarkStart', 'w:bookmarkEnd'):
        ids = [el.get(qn('w:id')) for el in d.element.body.iter(qn(tag))]
        if len(ids) != len(set(ids)):
            fails.append(f'id {tag} ganda (docx tidak valid untuk Word)')
    # identitas, status, tanda tangan
    ident = {r.cells[0].text.strip(): r.cells[1].text.strip() for r in d.tables[0].rows if len(r.cells) > 1}
    if ident.get('Document Status') != 'Draft':
        fails.append(f"Document Status harus Draft sampai sah, kini '{ident.get('Document Status')}'")
    if ident.get('Effective Date'):
        fails.append('Effective Date harus kosong sampai penyetuju menandatangani')
    for k in ('Document Title', 'Document Number', 'Department', 'Classification', 'Process Owner', 'Review Period'):
        if not ident.get(k):
            fails.append(f'Identity kosong: {k}')
    cls = (ident.get('Classification') or '').upper()
    if not cls.startswith(('PUBLIC', 'INTERNAL', 'RESTRICTED', 'SECRET')):
        fails.append(f"Classification harus tingkat SNX-ISMS-PRO-01 (PUBLIC/INTERNAL/RESTRICTED/SECRET), kini '{ident.get('Classification')}'")
    if cls.startswith(('RESTRICTED', 'SECRET')):
        foot = ' '.join(t.text or '' for s in d.sections for t in s.footer._element.iter(qn('w:t')))
        if cls.split()[0] not in foot.upper():
            fails.append('dokumen RESTRICTED/SECRET tanpa label di footer (PRO-01 Bagian 5)')
    cover = para_text(d.element.body[1])
    if ident.get('Document Number') and ident['Document Number'] not in cover:
        fails.append('Nomor di sampul tidak sama dengan Document Number')
    for r in d.tables[2].rows[1:]:
        if r.cells[3].text.strip():
            fails.append('kolom Signature harus kosong (diisi penyetuju sendiri)')
        if not r.cells[2].text.strip():
            fails.append(f'Approval kosong: {r.cells[1].text.strip()}')
    rev = d.tables[1].rows[1].cells
    if not rev[2].text.strip():
        fails.append('Revision History baris 00 tanpa deskripsi')

    body = list(d.element.body)
    heads = [(i, para_text(el)) for i, el in enumerate(body)
             if el.tag == qn('w:p') and el.find('.//' + qn('w:pStyle')) is not None
             and el.find('.//' + qn('w:pStyle')).get(qn('w:val')) == 'Heading1' and para_text(el)]
    names = [h for _, h in heads if h not in ('Identity Document', 'Revision History', 'Dokument Approval')
             and not h.startswith('Dokumen ')]
    if mode == 'sop':
        if names != SOP_HEADINGS:
            fails.append(f'urutan 11 heading tidak baku: {names}')
        # isi tiap bagian tidak kosong
        idx = {h: i for i, h in heads}
        order = [idx.get(h) for h in SOP_HEADINGS]
        for k, h in enumerate(SOP_HEADINGS):
            if order[k] is None:
                continue
            end = order[k + 1] if k + 1 < len(order) and order[k + 1] else len(body) - 1
            seg = body[order[k] + 1:end]
            has_text = any(para_text(e) for e in seg)
            if not has_text:
                fails.append(f'bagian kosong: {h}')
            if h == 'PROCESS FLOW' and not any(e.findall('.//' + qn('w:drawing')) for e in seg):
                fails.append('PROCESS FLOW tanpa diagram')
            if h == 'STANDARD OPERATING PROCEDURE':
                h2 = [e for e in seg if e.tag == qn('w:p') and e.find('.//' + qn('w:pStyle')) is not None
                      and e.find('.//' + qn('w:pStyle')).get(qn('w:val')) == 'Heading2']
                if not h2:
                    fails.append('bagian 6 tanpa sub-proses 6.x')
        # tabel isi tanpa sel kosong (kecuali tabel identitas, revisi, approval)
        for ti, t in enumerate(d.tables[3:], start=3):
            for ri, r in enumerate(t.rows[1:], start=1):
                empties = [c for c in r.cells if not c.text.strip()]
                if empties:
                    fails.append(f'tabel {ti} baris {ri}: {len(empties)} sel kosong')
                    break
        # PIC SLA harus ada di tabel peran
        roles = {r.cells[0].text.strip().lower() for r in d.tables[4].rows[1:]}
        sla_t = [t for t in d.tables if t.rows[0].cells[0].text.strip() == 'Activity']
        if sla_t:
            for r in sla_t[0].rows[1:]:
                pic = r.cells[1].text.strip().lower()
                if pic and not any(p.strip() in roles for p in re.split(r'[,/;]| dan ', pic)):
                    warns.append(f'PIC SLA "{r.cells[1].text.strip()}" tidak ada di tabel peran')
    else:
        if not names:
            fails.append('dokumen tanpa bagian isi')
    usulan = len(re.findall(r'\(USULAN\)', txt))
    if usulan:
        warns.append(f'{usulan} butir (USULAN) menunggu ACC')
    out = {'file': path, 'mode': mode, 'pass': not fails, 'fails': fails, 'warns': warns}
    print(json.dumps(out, ensure_ascii=False, indent=1))
    sys.exit(0 if not fails else 1)


if __name__ == '__main__':
    main()
