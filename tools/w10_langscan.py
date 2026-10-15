#!/usr/bin/env python3
"""Gerbang W10 (LEAD G-4): ukur bahasa dan kerapian form, plus cek substansi beku.

Pakai:
  python3 -I tools/w10_langscan.py selftest
  python3 -I tools/w10_langscan.py scan <berkas...>              -> JSON per dokumen
  python3 -I tools/w10_langscan.py diff <lama.docx|xlsx> <baru>  -> selisih substansi
"""
import json, re, sys, io, zipfile, html
import docx, openpyxl
from docx.oxml.ns import qn

ACTORS = [
    'Direktur Utama', 'Direktur', 'VP Commercial', 'VP Tech Operation and Services', 'Manager of Product Tech',
    'Brand & Marketing', 'Business Development', 'Legal & Compliance', 'Line Manager Legal', 'Staff Legal',
    'Finance Lead', 'Finance Staff', 'Finance', 'Accounting & Tax', 'HRGA Officer', 'HR Lead', 'Office Helper',
    'Technical/Operation Engineer L1', 'Head of Operation', 'Lead NOC', 'Lead SOC', 'NOC Engineer', 'SOC Analyst',
    'Service Delivery Manager', 'Project Manager', 'Project Coordinator', 'ISMS Coordinator', 'Admin Sistem',
    'Pemilik tindakan', 'Pemverifikasi', 'Pemeriksa', 'Pencatat', 'Penanggap alternatif', 'Penanggung Jawab',
    'Atasan Langsung Tujuan', 'Atasan Langsung Asal', 'Atasan Langsung', 'Unit Kerja Pemohon', 'Pejabat Berwenang',
    'Pemohon', 'Penyetuju', 'Approver', 'Reviewer', 'Requester', 'PIC', 'Management', 'Commercial', 'Operation',
    'FAT', 'Legal', 'HR', 'SINUX', 'klien', 'customer', 'pelanggan', 'tim', 'karyawan', 'vendor', 'pemasok',
]
PERSONS = ['Samuel Kristanto Utomo', 'Amirull Yusuf Zulkarnain', 'Laras Rasmawati', 'Fani Fatullah', "Ahmad Rifa'i",
           'Zilly Dwinata', 'Rolin Krisnawati Nababan', 'Ester Febe Utomo', 'Alif Altamziz',
           'Kana Salsabila Putri Koeswologito', 'Pradipta Nugraha', 'Sita Nadia Salsabila Widodo', 'Erik Dwi Yanto',
           'Aziz Priyo Widakdo', 'Muhamad Reyhan Bahtiar']
INSTR = re.compile(r'^(PETUNJUK PENGISIAN|CARA MENGISI|KETENTUAN PENGISIAN|CARA PAKAI)\b', re.I)
PREFIX = r'(?:me|di|ber)'
_act = '|'.join(re.escape(a) for a in sorted(ACTORS, key=len, reverse=True))
CALQUE = re.compile(r'\byang (?:' + _act + r') (?!' + PREFIX + r')(?!adalah\b|sudah\b|belum\b|akan\b|dapat\b|harus\b|tidak\b|wajib\b|juga\b|masih\b|hanya\b|telah\b|sedang\b|lain\b|ini\b|itu\b|baru\b|lama\b|sama\b)([a-z]{3,})\b')
LEWAT = re.compile(r'\blewat\b', re.I)
POINTER = re.compile(r'\b(?:ada|diatur|dirinci|dijelaskan|tercantum|dimuat) (?:di|pada|dalam) (?:butir|bagian) \d', re.I)
USUL = re.compile(r'\(USULAN')
STARTS_ROLE = re.compile(r'^\s*(?:\d+[.)]\s*)?(?:' + _act + r')\b')
BUTIR = re.compile(r'^\s*(\d+(?:\.\d+)+)\s')
NUM = re.compile(r'(?<![\w.])\d+(?:[.,]\d+)?%?(?![\w.]*\.\d)')
SNX = re.compile(r'SNX-[A-Z]+-[A-Z]+-\d\d')


def docx_units(path_or_blob):
    d = docx.Document(path_or_blob if not isinstance(path_or_blob, bytes) else io.BytesIO(path_or_blob))
    body = d.element.body
    units, head = [], ''
    for el in body.iterchildren():
        if el.tag == qn('w:p'):
            p = docx.text.paragraph.Paragraph(el, d)
            t = p.text.strip()
            if p.style.name == 'Heading 1' and t:
                head = t
            units.append(dict(t=t, head=head, cell=False, style=p.style.name, jc=_jc(el)))
        elif el.tag == qn('w:tbl'):
            for tc in el.iter(qn('w:tc')):
                for pe in tc.iter(qn('w:p')):
                    t = ''.join(x.text or '' for x in pe.iter(qn('w:t'))).strip()
                    units.append(dict(t=t, head=head, cell=True, style='', jc=_jc(pe)))
    tables = []
    for tbl in body.iter(qn('w:tbl')):
        rows = tbl.findall(qn('w:tr'))
        trpr = rows[0].find(qn('w:trPr')) if rows else None
        tables.append(dict(rows=len(rows), header=trpr is not None and trpr.find(qn('w:tblHeader')) is not None))
    return units, tables


def _jc(pe):
    ppr = pe.find(qn('w:pPr'))
    j = ppr.find(qn('w:jc')) if ppr is not None else None
    return j.get(qn('w:val')) if j is not None else None


def xlsx_units(path_or_blob):
    wb = openpyxl.load_workbook(path_or_blob if not isinstance(path_or_blob, bytes) else io.BytesIO(path_or_blob))
    units, sheets = [], []
    for ws in wb:
        head = ''
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.strip():
                    t = c.value.strip()
                    if INSTR.match(t):
                        head = t.upper()
                    units.append(dict(t=t, head=head if ws.title == 'Kontrol Dokumen' else ws.title, cell=True, style='', jc=None))
        ps = ws.page_setup
        sheets.append(dict(sheet=ws.title, freeze=ws.freeze_panes, fit_w=ws.sheet_properties.pageSetUpPr.fitToPage if ws.sheet_properties.pageSetUpPr else None,
                           fitToWidth=ps.fitToWidth, orientation=ps.orientation, print_title_rows=ws.print_title_rows))
    return units, sheets


def metrics(units):
    text = '\n'.join(u['t'] for u in units)
    instr = [u['t'] for u in units if INSTR.match(u['head'] or '') and u['t'] and not INSTR.match(u['t'])]
    role_start = [t for t in instr if STARTS_ROLE.match(t)]
    return dict(
        words=len(text.split()),
        calque=[m.group(0) for m in CALQUE.finditer(text)],
        lewat=len(LEWAT.findall(text)),
        pointer=len(POINTER.findall(text)),
        usulan=len(USUL.findall(text)),
        instr_items=len(instr), instr_role_start=len(role_start),
        justified_cells=sum(1 for u in units if u['cell'] and u['jc'] == 'both' and u['t']),
    )


def scan(path):
    if path.endswith('.docx'):
        units, tables = docx_units(path)
        m = metrics(units)
        m['tables_ge3_no_header'] = sum(1 for t in tables if t['rows'] >= 3 and not t['header'])
    else:
        units, sheets = xlsx_units(path)
        m = metrics(units)
        m['sheets'] = sheets
    return m


def facts(units):
    text = '\n'.join(u['t'] for u in units)
    butir = [BUTIR.match(u['t']).group(1) for u in units if BUTIR.match(u['t']) and not u['cell']]
    usul_butir = sorted({BUTIR.match(u['t']).group(1) for u in units if BUTIR.match(u['t']) and not u['cell'] and USUL.search(u['t'])})
    heads = [u['t'] for u in units if u['style'].startswith('Heading') and u['t']]
    return dict(butir=butir, usul_butir=usul_butir, headings=heads,
                snx=sorted(set(SNX.findall(text))), persons=sorted({p for p in PERSONS if p in text}),
                numbers=sorted(set(NUM.findall(text))))


def diff(old, new):
    get = docx_units if old.endswith('.docx') else xlsx_units
    a, b = facts(get(old)[0]), facts(get(new)[0])
    out = {}
    for k in a:
        if a[k] != b[k]:
            if isinstance(a[k], list) and k in ('snx', 'persons', 'numbers', 'usul_butir'):
                out[k] = dict(hilang=sorted(set(a[k]) - set(b[k])), baru=sorted(set(b[k]) - set(a[k])))
            else:
                out[k] = 'URUTAN/JUMLAH BERBEDA'
    return out


def selftest():
    pos = ['bukti pemungutan PPN yang SINUX terbitkan lewat CoreTax', 'Pajak yang klien potong dari pembayaran',
           'pihak yang Finance Lead tentukan']
    neg = ['yang diterbitkan SINUX melalui CoreTax', 'yang klien sudah isi', 'Overdue: invoice yang melewati jatuh tempo',
           'yang SINUX menerbitkan']
    for s in pos:
        assert CALQUE.search(s), 'calque pos ' + s
    for s in neg:
        assert not CALQUE.search(s), 'calque neg ' + s
    assert LEWAT.search('dikirim lewat email') and not LEWAT.search('invoice yang melewati jatuh tempo') and not LEWAT.search('terlewat')
    assert POINTER.search('Lama dan titik awal hitungannya ada di butir 4.8.') and not POINTER.search('sesuai butir 4.8')
    assert STARTS_ROLE.match('7. Lead NOC atau Lead SOC menambah baris') and not STARTS_ROLE.match('Tambah baris di tabel Tindakan')
    assert INSTR.match('PETUNJUK PENGISIAN') and INSTR.match('Cara mengisi') and not INSTR.match('ISI FORMULIR')
    u = [dict(t='4.15 Penetapan pelaku (USULAN).', cell=False, style='', head='', jc=None), dict(t='Nilai 14 hari', cell=False, style='', head='', jc=None)]
    f = facts(u)
    assert f['butir'] == ['4.15'] and f['usul_butir'] == ['4.15'] and '14' in f['numbers']
    print('OK selftest w10_langscan')


if __name__ == '__main__':
    cmd = sys.argv[1]
    if cmd == 'selftest':
        selftest()
    elif cmd == 'scan':
        print(json.dumps({p: scan(p) for p in sys.argv[2:]}, ensure_ascii=False, indent=1))
    elif cmd == 'diff':
        print(json.dumps(diff(sys.argv[2], sys.argv[3]), ensure_ascii=False, indent=1))
